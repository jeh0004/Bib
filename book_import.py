"""Admin-only XLSX book import. Each source inventory number represents one physical copy."""
import io
import json
import secrets
from flask import Blueprint, request, render_template_string, session, redirect, url_for, flash
from openpyxl import load_workbook

bp = Blueprint("book_import", __name__)
HEADERS = ["Buchnr", "Autor", "Titel", "Medientyp", "Verlag", "Gebietsthema",
           "Sachthema", "Auflagedatum", "Buchindex", "ISBN", "Kiste",
           "Ausleihe Start", "Ausleihe Ende", "Ausleihe User"]
PAGE = """
{% extends 'base.html' %}
{% block title %}Excel-Buchimport{% endblock %}
{% block content %}
<h1>Excel-Buchimport</h1>
<p>DAVLU-Datei (.xlsx) hochladen. Zuerst <strong>nur prüfen</strong>, anschließend bewusst importieren.
Bestehende Inventarnummern werden übersprungen. Bestehende Bücher und Ausleihen werden nicht verändert.</p>
<form method="post" enctype="multipart/form-data">
<input type="hidden" name="csrf_token" value="{{ session.csrf_token }}">
<div class="mb-3"><input class="form-control" type="file" name="file" accept=".xlsx" required></div>
<button class="btn btn-outline-primary" type="submit" name="mode" value="preview">Datei prüfen</button>
<button class="btn btn-primary" type="submit" name="mode" value="import"
 onclick="return confirm('Neue Bücher und Exemplare wirklich importieren?')">Jetzt importieren</button>
</form>
{% if result %}
<div class="alert alert-{{ 'danger' if result.errors else 'info' }} mt-3">
Zeilen: {{ result.rows }} · Neue Exemplare: {{ result.new }} · Bereits vorhanden: {{ result.skipped }}
· Fehler: {{ result.errors|length }}.
{% if result.saved %}<strong>Import gespeichert.</strong>{% else %}<strong>Keine Änderungen gespeichert.</strong>{% endif %}
</div>
{% if result.errors %}<ul>{% for e in result.errors[:20] %}<li>{{ e }}</li>{% endfor %}</ul>{% endif %}
{% endif %}
{% endblock %}
"""

def value(x):
    if x is None:
        return ""
    return str(x).strip()

def inventory(x):
    if isinstance(x, float) and x.is_integer():
        x = int(x)
    return value(x)

def parse_xlsx(data):
    if not data[:2] == b"PK":
        raise ValueError("Ungültige XLSX-Datei")
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        head = [value(x) for x in next(rows)]
        if head[:len(HEADERS)] != HEADERS:
            raise ValueError("Spalten stimmen nicht mit der DAVLU-Vorlage überein")
        entries = []
        errors = []
        seen = set()
        for n, raw in enumerate(rows, 2):
            if not any(v is not None and value(v) for v in raw):
                continue
            cells = list(raw[:len(HEADERS)])
            cells += [None] * (len(HEADERS) - len(cells))
            row = dict(zip(HEADERS, cells))
            inv = inventory(row["Buchnr"])
            if not inv or not value(row["Titel"]):
                errors.append(f"Zeile {n}: Buchnr oder Titel fehlt")
            elif inv in seen:
                errors.append(f"Zeile {n}: doppelte Buchnummer {inv}")
            elif any(value(row[x]) for x in ("Ausleihe Start", "Ausleihe Ende", "Ausleihe User")):
                errors.append(f"Zeile {n}: Ausleihdaten vorhanden; manuelle Prüfung erforderlich")
            else:
                seen.add(inv)
                entries.append((inv, row))
        return entries, errors
    finally:
        wb.close()

def import_rows(db, entries, errors, commit=False):
    existing = {r[0] for r in db.execute(
        "SELECT inventory_code FROM book_copies WHERE inventory_code IS NOT NULL")}
    new = [(inv, row) for inv, row in entries if inv not in existing]
    result = dict(rows=len(entries)+len(errors), new=len(new),
                  skipped=len(entries)-len(new), errors=list(errors), saved=False)
    if errors or not commit:
        return result
    # A single transaction: either every new record is committed or none.
    with db:
        for inv, row in new:
            metadata = {key: value(row[key]) for key in (
                "Verlag", "Gebietsthema", "Sachthema", "Auflagedatum", "Buchindex")}
            title = value(row["Titel"])
            author = value(row["Autor"]) or "Unbekannt"
            category = value(row["Medientyp"])
            isbn = value(row["ISBN"])
            cur = db.execute(
                "INSERT INTO books(title,author,isbn,category,description,total_copies) VALUES (?,?,?,?,?,1)",
                (title, author, isbn or None, category or None,
                 json.dumps(metadata, ensure_ascii=False)))
            book_id = cur.lastrowid
            db.execute("""INSERT INTO book_copies
                (book_id,copy_number,inventory_code,qr_token,location,notes)
                VALUES (?,1,?,?,?,?)""",
                (book_id, inv, secrets.token_urlsafe(18),
                 ("Kiste " + value(row["Kiste"])) if value(row["Kiste"]) else None,
                 "DAVLU-Buchnummer: " + inv))
    result["saved"] = True
    return result

@bp.route("/admin/books/import", methods=["GET", "POST"])
def upload():
    from app import get_db
    if not session.get("user_id"):
        return redirect(url_for("login"))
    if session.get("role") != "admin":
        return ("Forbidden", 403)
    result = None
    if request.method == "POST":
        f = request.files.get("file")
        if not f or not f.filename.lower().endswith(".xlsx"):
            flash("Bitte eine XLSX-Datei auswählen", "danger")
        else:
            try:
                entries, errors = parse_xlsx(f.read())
                result = import_rows(get_db(), entries, errors,
                                     commit=request.form.get("mode") == "import")
            except (ValueError, OSError, IndexError, StopIteration) as exc:
                flash("Import fehlgeschlagen: " + str(exc), "danger")
    return render_template_string(PAGE, result=result)
